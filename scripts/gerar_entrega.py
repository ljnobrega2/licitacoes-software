from triagem import *
from collections import Counter, defaultdict
from decimal import Decimal, ROUND_DOWN
import csv
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

MANUAL={
'43976166000150-1-000089/2026':'Edital/TR lidos: 30 assentos Standard por 12 meses; exclusivo ME/EPP; ativação em até 10 dias úteis; comprovar fornecimento legítimo. O edital admite documento comercial idôneo e dispensa declaração específica de revendedor autorizado quando não fizer parte do modelo do fabricante. Plataforma BNC. É a contratação da primeira imagem enviada.',
'83102228000110-1-000108/2026':'TR lido: 6 acessos institucionais ChatGPT Business por 12 meses, pagamento anual, configuração e suporte de licenciamento. Não inclui desenvolvimento, infraestrutura local ou integração por API.',
'07842278000155-1-000013/2026':'Edital lido: 12 licenças ChatGPT Business por 12 meses; intermediação de pagamento, suporte e gerenciamento. Portal de Compras Públicas; sessão em 07/10/2026 às 09:01; cadastro PNCP encerra recebimento às 09:00.',
'78316064000193-1-000033/2026':'Edital lido: contratação de broker; grupos Claude Standard/Premium, créditos/assentos adicionais e ChatGPT Business. Itens: 36 Standard Claude, 1 Premium Claude, provisão de créditos e 3 ChatGPT. Sessão 16/10 às 13:30. Possível publicação duplicada do controle 32.',
'78316064000193-1-000032/2026':'Possível duplicidade do controle 33: mesmo órgão, objeto, prazo e valor. Conferir edital oficial; não contar como segunda oportunidade independente.',
'11552755000115-1-000047/2026':'Objeto informa 2 usuários e 12 meses. Nome comercial ChatGPT-0.5 no cadastro deve ser esclarecido no TR. Aviso disponibilizado como ZIP; conteúdo não integralmente extraído.',
'88488366000100-1-000400/2026':'Aviso lido: assinatura anual Canva Pro+ Equipes para 6 usuários; item único (quantidade 1 refere-se ao pacote de 6 usuários).',
'24134488000108-1-000155/2026':'Divergência a esclarecer: objeto menciona VIP governamental, mas itens especificam EDUCACIONAL NAMED; 7 Creative Cloud e 5 Acrobat, por 12 meses. Não cotar plano varejista sem conferir o TR.',
'82804212000196-1-000324/2026':'Edital retificado lido: 2 subscrições anuais Adobe Acrobat Pro. Possível duplicidade do controle 302; priorizar leitura da versão retificada.',
'82804212000196-1-000302/2026':'Possível duplicidade do controle 324; verificar edital retificado e situação no portal de disputa.',
'62759595000110-1-000041/2026':'Aviso/TR lidos: 4 Adobe Creative Cloud VIP Teams por 24 meses, vinculados à organização institucional da ANSN; canal legítimo e suporte; não aceita licenças educacionais/trial/OEM/usadas. Lances 08/10 das 08h às 14h; proposta antes das 08h.',
'00394452000103-1-021390/2026':'Microsoft 365 Business Basic. Aviso baixado, PDF sem texto extraível; condições complementares do cadastro indicam preferência ME/EPP e entrega em 30 dias.',
'34164319000174-1-000189/2026':'O próprio objeto exige comprovação de revenda autorizada Adobe com especialização em governo.'
,'00394452000103-1-021427/2026':'Divergência a esclarecer: objeto descreve software de orçamentação para manutenção de frota, mas a descrição do item no catálogo menciona antivírus corporativo. Não assumir que se trata de compra de antivírus.'
}

def ceiling(v):
    return float((Decimal(str(v))*Decimal('0.79')).quantize(Decimal('0.01'),rounding=ROUND_DOWN)) if v is not None and v>0 else None

def dtvalue(v):
    try:return dt.datetime.fromisoformat(v)
    except:return v or ''

def clean(v):
    if not isinstance(v,str):return v
    v=ILLEGAL_CHARACTERS_RE.sub('',v)
    return "'"+v if v.startswith(('=','+','-','@')) else v

def publiclink(r):
    u=r.get('item_url') or ''
    if u.startswith('/compras/'):return 'https://pncp.gov.br/app/editais/'+u.removeprefix('/compras/')
    return 'https://pncp.gov.br/app'+u

def main():
    records=load_rows()
    details={}
    for p in (DATA/'detalhes').glob('*.json'):
        d=json.loads(p.read_text());details[d['id']]=d
        # Resposta vazia (204) da API é guardada como objeto; itens e arquivos precisam ser listas.
        for k in ('itens','arquivos'):
            if k in d and not isinstance(d[k],list):d[k]=d[k].get('data',[]) if isinstance(d[k].get('data'),list) else []
    for r in records:
        det=details.get(r['numero_controle_pncp'],{})
        comp=det.get('compra',{})
        if comp:
            for dest,src in [('description','objetoCompra'),('informacao_complementar','informacaoComplementar'),('data_inicio_vigencia','dataAberturaProposta'),('data_fim_vigencia','dataEncerramentoProposta'),('valor_total_estimado','valorTotalEstimado'),('situacao_id','situacaoCompraId'),('situacao_nome','situacaoCompraNome'),('link_sistema_origem','linkSistemaOrigem'),('link_processo_eletronico','linkProcessoEletronico')]:
                r[dest]=comp.get(src)
        extra=' '.join(i.get('descricao','') for i in det.get('itens',[]) if isinstance(i,dict))
        r['_classe'],r['_marcas']=classify(r,extra)
        if r['numero_controle_pncp']=='00394452000103-1-021427/2026':
            r['_classe']='Software sem marca / validar fornecimento';r['_marcas']=[]
        r['_status']=status(r)
        r['_detalhes']=det
        r['_link']=publiclink(r)
    # Preserve duplicates in an audit sheet, without double-counting the shortlist.
    groups=defaultdict(list)
    for r in records:
        if r['document_type']!='edital':continue
        desc=re.sub(r'^\[[^\]]+\]\s*-?\s*','',norm(r.get('description') or ''))
        key=(r['orgao_cnpj'],re.sub(r'\W+','',desc),r.get('data_fim_vigencia'),r.get('valor_total_estimado'))
        groups[key].append(r)
    for group in groups.values():
        if len(group)>1:
            group.sort(key=lambda r:(len(r['_detalhes']),int(r['numero_sequencial'])),reverse=True)
            for r in group[1:]:r['_duplicado_de']=group[0]['numero_controle_pncp']
    ativos={'Prazo aberto no cadastro PNCP','Recebimento futuro anunciado'}
    buckets=defaultdict(list)
    for r in records:
        cl=r['_classe'];st=r['_status']
        if cl.startswith(('Fora','Curso')):bucket='Revisar resultados'
        elif r.get('_duplicado_de'):bucket='Duplicidades possíveis'
        elif r['document_type']=='ata':bucket='Atas vigentes' if st in ['Ata vigente — fornecedor já selecionado','Ata com vigência futura'] else 'Vencidas e inconsistentes'
        elif st not in ativos:bucket='Vencidas e inconsistentes'
        elif cl=='Licenças comerciais identificadas':bucket='Licenças comerciais'
        elif cl=='Software sem marca / validar fornecimento':bucket='Outros softwares'
        else:bucket='Soluções com serviços'
        buckets[bucket].append(r)
    for rows in buckets.values():rows.sort(key=lambda r:(r.get('data_fim_vigencia') or '9999',r.get('uf',''),r.get('municipio_nome','')))
    headings=['Situação apurada','Enquadramento','Produtos/marcas','Órgão','Unidade compradora','Esfera','Cidade','UF','Edital / aviso','Prazo de propostas / fim vigência','Início propostas / vigência','Valor estimado TOTAL (R$)','Teto de custo TOTAL p/ margem 15% (R$)','Modalidade','Registro de preços','Objeto integral','Pontos a conferir','Verificação realizada','Link PNCP','Portal de disputa / origem','Edital ou aviso (arquivo)','Termo de referência (arquivo)','Controle PNCP','Duplicidade provável de','Publicação PNCP','Adesão da ata (cadastro)']
    def row(r):
        d=r['_detalhes'];files=[f for f in d.get('arquivos',[]) if f.get('statusAtivo')]
        ed=next((f['url'] for f in files if f.get('tipoDocumentoId') in [1,2]),files[0]['url'] if files else '')
        tr=next((f['url'] for f in files if f.get('tipoDocumentoId')==4),'')
        note=MANUAL.get(r['numero_controle_pncp'],'')
        if not note:
            if r['document_type']=='ata':note='Fornecedor já selecionado; serve para inteligência comercial. Não é novo prazo de disputa. Adesão não transforma terceiro em fornecedor da ata.'
            elif r['_classe']=='Licenças comerciais identificadas':note='Candidata a fornecimento por canal/revenda. Confirmar SKU, prazo, elegibilidade do cliente, credenciamento do fabricante, suporte e habilitação no edital.'
            elif r['_classe']=='Software sem marca / validar fornecimento':note='Marca/modelo ou modalidade não definidos na triagem. Conferir itens e TR; possibilidade de revenda não confirmada.'
            else:note='Pode envolver implantação, operação, hardware ou desenvolvimento. Não presumir fornecimento apenas de assinatura.'
        if r.get('_duplicado_de'):note+=' Possível duplicidade de '+r['_duplicado_de']+'.'
        ver='Cadastro de busca oficial PNCP' if not r.get('_fonte') else 'API oficial PNCP de propostas abertas (varredura por UF)'
        if 'itens' in d:ver+=' + itens oficiais'
        if 'compra' in d:ver+=' + detalhe da contratação'
        if r['numero_controle_pncp'] in MANUAL:ver+='; observação específica na coluna anterior'
        val=r.get('valor_total_estimado') if r['document_type']=='edital' else r.get('valor_global')
        if not val or r.get('indicador_orcamento_sigiloso'):val=None
        adh=r.get('permite_adesao')
        return [r['_status'],r['_classe'],'; '.join(r['_marcas']),r.get('orgao_nome'),r.get('unidade_nome'),r.get('esfera_nome'),r.get('municipio_nome'),r.get('uf'),r.get('title'),dtvalue(r.get('data_fim_vigencia')),dtvalue(r.get('data_inicio_vigencia')),val,ceiling(val),r.get('modalidade_licitacao_nome'),'Sim' if r.get('srp') else 'Não/não informado',r.get('description'),note,ver,r['_link'],r.get('link_sistema_origem') or '',ed,tr,r['numero_controle_pncp'],r.get('_duplicado_de',''),dtvalue(r.get('data_publicacao_pncp')),'Sim' if adh is True else 'Não' if adh is False else 'Não informado']

    wb=Workbook();summary=wb.active;summary.title='Leia primeiro'
    summary.append(['LEVANTAMENTO NACIONAL — LICENÇAS E ASSINATURAS DE SOFTWARE',''])
    summary.append(['Consulta concluída',NOW.strftime('%d/%m/%Y %H:%M')+' — horário de Brasília'])
    summary.append(['Abrangência','Busca nacional, sem filtro de UF, município, esfera, porte ou valor; PNCP e checagem complementar de páginas oficiais.'])
    summary.append(['Limite de cobertura','LEVANTAMENTO NÃO EXAUSTIVO: houve falhas em páginas da API; cadastros/portais não integrados e itens com descrições genéricas podem não ter sido encontrados. Não comprova a ausência de outras oportunidades.'])
    summary.append(['O que significa candidata','A classificação identifica objeto comercial compatível com licenciamento; não comprova que sua empresa está habilitada nem autorizada pelo fabricante a revendê-lo.'])
    summary.append(['Datas','Prazos extraídos do cadastro PNCP; conferir eventos, suspensões, retificações e horário no portal da disputa. Datas de atas são vigência, não prazo para ofertar.'])
    summary.append(['Valores','Valor estimado global da contratação, não lucro nem preço efetivo de compra. Zero/ausente/sigiloso fica em branco. Valores por licença estão na aba Itens para cotar quando disponíveis.'])
    summary.append(['Fórmula solicitada anteriormente','Teto de custo = referência × 0,79, considerando imposto 6% e margem 15% sobre venda; custo inclui licenças e TODAS as despesas. Não é lance mínimo. Lance mínimo = custo efetivo / 0,79.'])
    summary.append(['Atas','Atas vigentes têm fornecedor selecionado; usar para conhecer compradores, preços e possíveis renovações, sem confundir com oportunidade aberta para novo fornecedor.'])
    summary.append(['Duplicidades','Publicações com mesmo CNPJ, descrição normalizada, prazo e valor são sinalizadas como possíveis duplicidades e preservadas em aba própria. Confirmar os processos antes de descartá-las.'])
    summary.append(['Verificação',str(len(details))+' contratações tiveram itens e lista de documentos consultados; leitura focal de 15 arquivos de 10 contratações prioritárias, sem auditoria integral de todos os editais. Alguns PDFs/ZIP não tiveram texto extraível.'])
    uf_audit=json.loads((DATA/'coleta_uf_resumo.json').read_text()) if (DATA/'coleta_uf_resumo.json').exists() else None
    if uf_audit:
        falhas=sum(len(a.get('falhas',[])) for a in uf_audit['ufs'])
        summary.append(['Varredura por UF','Além das buscas por termo, a API oficial de propostas abertas foi percorrida UF por UF: '+str(uf_audit['unicos'])+' contratações com recebimento de propostas aberto lidas, '+str(uf_audit['candidatasObjeto'])+' pré-selecionadas pelo objeto para triagem. Páginas que falharam: '+str(falhas)+'. A triagem lê o OBJETO; software citado apenas em item ou anexo pode não aparecer.'])
    summary.append(['Atualização','Esta é uma fotografia da consulta. Não foi ativado monitoramento automático.'])
    summary.append(['',''])
    for name in ['Licenças comerciais','Outros softwares','Soluções com serviços','Atas vigentes','Duplicidades possíveis','Vencidas e inconsistentes','Revisar resultados']:summary.append([name,len(buckets[name])])
    summary.append(['Registros PNCP únicos recuperados',len(records)])
    active=sum((buckets[x] for x in ['Licenças comerciais','Outros softwares','Soluções com serviços']),[])
    summary.append(['UFs com candidatos ativos encontrados',', '.join(sorted({r['uf'] for r in active}))])
    summary.append(['Cidade/UF','Localidade do órgão comprador. Requisitos de prestação presencial/regionalidade dependem do edital.'])
    summary.append(['Documentação PNCP','https://pncp.gov.br/api/consulta/swagger-ui/index.html'])
    summary.append(['Página oficial AGEHAB complementar','https://goias.gov.br/agehab/pregao-eletronico-no-006-2026/'])
    summary.column_dimensions['A'].width=42;summary.column_dimensions['B'].width=120
    for cells in summary:
        for c in cells:c.alignment=Alignment(vertical='top',wrap_text=True)
        summary.row_dimensions[cells[0].row].height=45 if cells[0].row<(14 if uf_audit else 13) else 30
    summary['A1'].font=Font(size=16,bold=True,color='FFFFFF');summary['A1'].fill=PatternFill('solid',fgColor='17365D');summary['B1'].fill=PatternFill('solid',fgColor='17365D')

    def sheet(name,headers,rows):
        ws=wb.create_sheet(name);ws.append(headers)
        for values in rows:
            ws.append([clean(v) for v in values])
        ws.freeze_panes='D2';ws.sheet_view.zoomScale=85
        if rows:
            t=Table(displayName='Tabela'+str(len(wb.worksheets)),ref=f'A1:{get_column_letter(len(headers))}{len(rows)+1}')
            t.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True);ws.add_table(t)
        for c in ws[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='17365D');c.alignment=Alignment(wrap_text=True,vertical='center')
        ws.row_dimensions[1].height=44
        for j,h in enumerate(headers,1):
            width=24
            if 'Objeto' in h or 'Descrição' in h or 'Pontos' in h:width=70
            elif 'Link' in h or 'arquivo' in h or 'Portal' in h:width=28
            elif h=='UF':width=7
            elif h in ['Órgão','Unidade compradora']:width=40
            elif 'Prazo' in h or 'Início' in h:width=23
            ws.column_dimensions[get_column_letter(j)].width=width
        for rr in ws.iter_rows(min_row=2):
            ws.row_dimensions[rr[0].row].height=54
            for c in rr:
                c.alignment=Alignment(vertical='top',wrap_text=True)
                if isinstance(c.value,dt.datetime):c.number_format='dd/mm/yyyy hh:mm'
                if isinstance(c.value,str) and c.value.startswith('https://'):
                    c.hyperlink=c.value;c.font=Font(color='0563C1',underline='single')
                if '(R$)' in headers[c.column-1] and isinstance(c.value,(float,int)):c.number_format='"R$" #,##0.00'
        return ws

    priority_ids=['83102228000110-1-000108/2026','07842278000155-1-000013/2026','00394452000103-1-021390/2026','07954480000179-1-025425/2026','43976166000150-1-000089/2026','62759595000110-1-000041/2026','88488366000100-1-000400/2026','24134488000108-1-000155/2026','11552755000115-1-000047/2026','78316064000193-1-000033/2026','82804212000196-1-000324/2026','00348003000110-1-001123/2026']
    priority=[r for r in active if r['numero_controle_pncp'] in priority_ids]
    priority.sort(key=lambda r:r.get('data_fim_vigencia',''))
    ph=['Prazo de propostas','Produtos','Cidade/UF','Órgão / unidade','Valor estimado TOTAL (R$)','Objeto','Pontos a conferir','Link oficial PNCP','Controle PNCP']
    sheet('Comece aqui',ph,[[dtvalue(r.get('data_fim_vigencia')),'; '.join(r['_marcas']),r.get('municipio_nome','')+'/'+r.get('uf',''),r.get('unidade_nome') or r.get('orgao_nome'),r.get('valor_total_estimado') or None,r.get('description'),MANUAL.get(r['numero_controle_pncp'],'Conferir edital, fornecedor e escopo.'),r['_link'],r['numero_controle_pncp']] for r in priority])
    for name in ['Licenças comerciais','Outros softwares','Soluções com serviços','Atas vigentes','Duplicidades possíveis','Vencidas e inconsistentes','Revisar resultados']:
        sheet(name,headings,[row(r) for r in buckets[name]])
    # Item-level quantities preserve the original procurement unit (not always one seat).
    ih=['Controle PNCP','Órgão / unidade','Cidade','UF','Prazo de propostas','Item','Descrição integral do item','Quantidade','Unidade','Referência UNITÁRIA (R$)','Referência TOTAL do item (R$)','Teto de custo UNITÁRIO p/ 15% (R$)','Cotação real UNITÁRIA com despesas (R$)','Lance mínimo UNITÁRIO calculado (R$)','Link PNCP','Observação']
    itemrows=[]
    for r in active:
        for item in r['_detalhes'].get('itens',[]):
            if not isinstance(item,dict):continue
            unit=item.get('valorUnitarioEstimado');total=item.get('valorTotal')
            if item.get('orcamentoSigiloso'):unit=total=None
            itemrows.append([r['numero_controle_pncp'],r.get('unidade_nome') or r.get('orgao_nome'),r.get('municipio_nome'),r.get('uf'),dtvalue(r.get('data_fim_vigencia')),item.get('numeroItem'),item.get('descricao'),item.get('quantidade'),item.get('unidadeMedida'),unit,total,ceiling(unit),None,None,r['_link'],'Conferir o período e se a unidade é usuário, pacote, mês ou crédito. Preencher cotação com todos os custos; fórmula considera imposto 6% e margem 15% sobre venda.'])
    ws=sheet('Itens para cotar',ih,itemrows)
    for i in range(2,len(itemrows)+2):
        ws[f'N{i}']=f'=IF(M{i}="","",ROUNDUP(M{i}/0.79,2))';ws[f'N{i}'].number_format='"R$" #,##0.00';ws[f'M{i}'].fill=PatternFill('solid',fgColor='FFF2CC')
    # One opportunity independently confirmed on an official state website, not found in the collected PNCP subset.
    eh=['Órgão','Cidade','UF','Objeto','Sessão/prazo','Valor','Identificador','Fonte oficial','Observação']
    er=[['AGEHAB — Agência Goiana de Habitação','Goiânia','GO','Autodesk AEC Collection, subscrição por 3 anos, treinamento básico e avançado',dt.datetime(2026,10,15,14,0),'Sigiloso','PE 006/2026; SISLOG 122608','https://goias.gov.br/agehab/pregao-eletronico-no-006-2026/','Aviso oficial atualizado em 02/10/2026; inclui transferência de conhecimento. Cadastro em SISLOG. Não confundir com revenda sem serviço.']]
    sheet('Fonte estadual complementar',eh,er)
    # Fonte externa: Portal de Compras Públicas. Resumo do objeto vem truncado pela API; conferir no link.
    pcprows=[]
    if (DATA/'externo_pcp.json').exists():
        bylink={};byorg=set();bycity={}
        for r in records:
            m=re.search(r'-(\d+)$',r.get('link_sistema_origem') or '')
            if m and 'portaldecompraspublicas' in r['link_sistema_origem']:bylink[m.group(1)]=r['numero_controle_pncp']
            if r['document_type']=='edital':byorg.add((norm(r.get('orgao_nome') or ''),(r.get('data_fim_vigencia') or '')[:16]))
            if r['document_type']=='edital':bycity[(norm(r.get('municipio_nome') or ''),r.get('uf'),(r.get('data_fim_vigencia') or '')[:16])]=r['numero_controle_pncp']
        for x in json.loads((DATA/'externo_pcp.json').read_text()):
            cl,brands=classify({'description':x.get('resumo') or ''})
            if cl.startswith(('Fora','Curso')):continue
            try:end=dt.datetime.fromisoformat(x['dataHoraFinalPropostas'].replace('Z','+00:00')).astimezone(ZoneInfo('America/Sao_Paulo')).replace(tzinfo=None)
            except Exception:end=None
            if end and end<=NOW:continue
            un=x.get('unidadeCompradora') or {}
            pncp=bylink.get(str(x['codigoLicitacao'])) or (end and bycity.get((norm(un.get('cidade') or ''),un.get('uf'),end.isoformat()[:16])) and 'Provável: '+bycity[(norm(un.get('cidade') or ''),un.get('uf'),end.isoformat()[:16])]+' (mesma cidade e prazo)') or ('Provável (mesmo órgão e prazo)' if end and (norm(x.get('razaoSocial') or ''),end.isoformat()[:16]) in byorg else 'Não localizado no conjunto PNCP coletado')
            pcprows.append([cl,'; '.join(brands),x.get('razaoSocial'),un.get('cidade'),un.get('uf'),((x.get('tipoLicitacao') or {}).get('siglaTipoLicitacao') or '')+' '+str(x.get('numero')),end or '',x.get('resumo'),'Sim' if x.get('isExclusivoME') else 'Não',(x.get('statusProcessoPublico') or {}).get('descricao'),x['_link'],pncp])
        pcprows.sort(key=lambda v:(v[0]!='Licenças comerciais identificadas',str(v[6])))
        sheet('Portal Compras Públicas',['Enquadramento','Produtos/marcas','Órgão','Cidade','UF','Processo','Prazo de propostas (Brasília)','Objeto (resumo truncado pelo portal)','Exclusivo ME/EPP','Situação no portal','Link do portal de disputa','Correspondência no PNCP'],pcprows)
    # Fonte externa: SEBRAE (Sistema S) não publica no PNCP. "Em andamento" inclui sessões já realizadas.
    sebrows=[]
    if (DATA/'externo_sebrae.json').exists():
        seb=json.loads((DATA/'externo_sebrae.json').read_text())
        for x in seb['licitacoes']:
            cl,brands=classify({'description':x.get('Objeto') or ''})
            if cl.startswith(('Fora','Curso')):continue
            ab=dtvalue(x.get('abertura'))
            st='Sessão futura — conferir edital no portal' if isinstance(ab,dt.datetime) and ab>NOW else 'Sessão já realizada; processo ainda em andamento'
            m=re.search(r'SEBRAE[-/ ]?([A-Z]{2})\b',x.get('Numero') or '')
            sebrows.append([st,cl,'; '.join(brands),'SEBRAE'+('/'+m.group(1) if m else ' Nacional'),m.group(1) if m else '',x.get('Numero'),x.get('ModalidadeDescricao'),ab,x.get('Objeto'),x.get('NumeroProcesso'),seb['fonte'],'Localizar pelo número no Canal do Fornecedor SEBRAE. Regulamento próprio do Sistema S (não é Lei 14.133). Disputa exige cadastro no portal.'])
        sebrows.sort(key=lambda v:(not v[0].startswith('Sessão futura'),str(v[7])))
        sheet('Sistema S - SEBRAE',['Situação apurada','Enquadramento','Produtos/marcas','Órgão','UF','Número','Modalidade','Abertura da sessão','Objeto integral','Processo','Fonte oficial','Observação'],sebrows)
    # Coverage including failures instead of claiming complete national coverage.
    audit=[]
    for filename,stage in [('busca_resumo.json','Busca inicial'),('recuperacao_resumo.json','Recuperação'),('lacunas_resumo.json','Lacunas refeitas com controle de taxa')]:
        if (DATA/filename).exists():
            for a in json.loads((DATA/filename).read_text()):
                err=a.get('erro') or ('; '.join('página '+str(e.get('pagina')) for e in a.get('falhas',[])))
                audit.append([stage,a.get('tipo'),a.get('termo'),a.get('totalDeclarado'),a.get('coletados'),'Parcial/falhou' if err else 'Paginação concluída na consulta',err or ''])
    if (DATA/'coleta_uf_resumo.json').exists():
        for a in json.loads((DATA/'coleta_uf_resumo.json').read_text())['ufs']:
            err='; '.join('página '+str(e.get('pagina')) for e in a.get('falhas',[]))
            audit.append(['Varredura da API de propostas abertas','edital','UF '+a['uf'],a.get('totalDeclarado'),a.get('coletados'),'Parcial/falhou' if err or a.get('totalDeclarado') is None else 'Todas as páginas da UF baixadas',err])
    if (DATA/'externo_pcp_resumo.json').exists():
        for t in json.loads((DATA/'externo_pcp_resumo.json').read_text())['termos']:
            audit.append(['Portal de Compras Públicas (fonte externa)','processo recebendo propostas',t['termo'],t.get('totalDeclarado'),t.get('coletados'),'Parcial/falhou' if t.get('erro') else 'Paginação concluída na consulta',t.get('erro') or ''])
    if (DATA/'externo_sebrae.json').exists():
        audit.append(['Canal do Fornecedor SEBRAE (fonte externa)','licitações em andamento','listagem nacional completa',seb['total'],seb['total'],'Listagem baixada e filtrada localmente',''])
    for nome,motivo in [('Licitações-e (Banco do Brasil)','HTTP 403 para acesso automatizado'),('Licitações CAIXA','página protegida por captcha'),('Petronect (Petrobras)','lista carregada por JavaScript; não extraída'),('SESI/SENAI/SESC/SENAC/SEST-SENAT','portais regionais separados; não percorridos'),('BLL, BNC, Licitanet, ComprasBR','busca pública por formulário; processos de órgãos públicos costumam constar no PNCP, mas isso não foi verificado um a um')]:
        audit.append(['Fonte externa NÃO coberta','—',nome,None,None,'Sem cobertura',motivo])
    sheet('Cobertura das buscas',['Etapa','Tipo','Termo','Total declarado','Coletados na etapa','Resultado','Falha'],audit)
    # Dados do painel HTML (painel/dados.js). Abas ativas levam itens e documentos; as demais vão resumidas.
    def iso(v):return v.isoformat() if isinstance(v,dt.datetime) else (v or None)
    def slim(r,full):
        v=row(r);d=r['_detalhes']
        o={'id':v[22],'st':v[0],'cl':v[1],'mc':r['_marcas'],'org':v[3],'un':v[4],'esf':v[5],'cid':v[6],'uf':v[7],'tit':v[8],'fim':iso(v[9]),'ini':iso(v[10]),'val':v[11],'mod':v[13],'srp':bool(r.get('srp')),'obj':clean_txt(v[15],4000 if full else 320),'link':v[18],'pub':iso(v[24])}
        if full:
            o.update({'nota':v[16],'ver':v[17],'orig':v[19] or None,'dup':v[23] or None,'manual':r['numero_controle_pncp'] in MANUAL,'info':clean_txt(r.get('informacao_complementar'),1500)})
            its=[i for i in d.get('itens',[]) if isinstance(i,dict)]
            o['it']=[{'n':i.get('numeroItem'),'d':clean_txt(i.get('descricao'),700),'q':i.get('quantidade'),'u':i.get('unidadeMedida'),'vu':None if i.get('orcamentoSigiloso') else i.get('valorUnitarioEstimado'),'vt':None if i.get('orcamentoSigiloso') else i.get('valorTotal'),'b':i.get('tipoBeneficioNome'),'s':i.get('situacaoCompraItemNome')} for i in its[:60]]
            o['nit']=len(its)
            o['me']=any('exclusiva' in norm(i.get('tipoBeneficioNome') or '') for i in its)
            o['docs']=[{'t':f.get('tipoDocumentoNome'),'n':f.get('titulo'),'u':f.get('url'),'p':f.get('dataPublicacaoPncp')} for f in d.get('arquivos',[]) if isinstance(f,dict) and f.get('statusAtivo')]
            o['det']='itens' in d
        elif r['document_type']=='ata':o['ad']=v[25]
        return o
    def clean_txt(t,n):
        t=ILLEGAL_CHARACTERS_RE.sub('',t or '').strip()
        return t if len(t)<=n else t[:n]+'…'
    painel={'geradoEm':NOW.isoformat(timespec='minutes'),
        'abertas':[slim(r,True) for n in ['Licenças comerciais','Outros softwares','Soluções com serviços','Duplicidades possíveis'] for r in buckets[n]],
        'atas':[slim(r,False) for r in buckets['Atas vigentes']],
        'vencidas':[slim(r,False) for r in buckets['Vencidas e inconsistentes']],
        'revisar':[slim(r,False) for r in buckets['Revisar resultados'] if r['document_type']=='edital' and r['_status'] in ativos],
        'pcp':[dict(zip(['cl','mc','org','cid','uf','proc','fim','obj','me','st','link','pncp'],[iso(x) for x in v])) for v in pcprows],
        'sebrae':[dict(zip(['st','cl','mc','org','uf','num','mod','ab','obj','proc','link','nota'],[iso(x) for x in v])) for v in sebrows],
        'estadual':[dict(zip(['org','cid','uf','obj','fim','val','id','link','nota'],[iso(x) for x in v])) for v in er],
        'cobertura':[dict(zip(['etapa','tipo','termo','total','col','res','falha'],v)) for v in audit],
        'varredura':uf_audit and {'unicos':uf_audit['unicos'],'candidatas':uf_audit['candidatasObjeto'],'em':uf_audit['consultadoEm'],'ufs':[{'uf':a['uf'],'total':a['totalDeclarado'],'col':a['coletados'],'falhas':len(a['falhas'])} for a in uf_audit['ufs']]},
        'totais':{'registros':len(records),'detalhes':len(details),'itens':len(itemrows)}}
    (ROOT/'painel').mkdir(exist_ok=True)
    (ROOT/'painel'/'dados.js').write_text('window.DADOS='+json.dumps(painel,ensure_ascii=False,separators=(',',':')).replace('</','<\\/').replace('\ufffd','?')+';')
    counts=Counter(r.get('uf') for r in active);spheres=Counter(r.get('esfera_nome') for r in active)
    sheet('Resumo por UF',['UF','Candidatos com prazo aberto/futuro'],[[uf,counts.get(uf,0)] for uf in ['AC','AL','AP','AM','BA','CE','DF','ES','GO','MA','MT','MS','MG','PA','PB','PR','PE','PI','RJ','RN','RS','RO','RR','SC','SP','SE','TO']])
    filepath=ROOT/'Licitacoes_software_Brasil_2026-10-06.xlsx';wb.save(filepath)
    csvdir=ROOT/'csv';csvdir.mkdir(exist_ok=True)
    for name,rows in buckets.items():
        slug=re.sub(r'[^a-z0-9]+','_',norm(name))
        with (csvdir/(slug+'.csv')).open('w',encoding='utf-8-sig',newline='') as f:
            w=csv.writer(f,delimiter=';');w.writerow(headings)
            for r in rows:w.writerow([v.isoformat(' ') if isinstance(v,dt.datetime) else clean(v) for v in row(r)])
    (DATA/'resultado_final.json').write_text(json.dumps([{k:v for k,v in r.items() if k!='_detalhes'} for r in records],ensure_ascii=False))
    stats={'consulta':NOW.isoformat(),'recuperados':len(records),'abas':{k:len(v) for k,v in buckets.items()},'itens':len(itemrows),'ufs':dict(counts),'esferas':dict(spheres),'detalhes':len(details),'sebraeSoftware':len(sebrows),'sebraeSessaoFutura':sum(1 for v in sebrows if v[0].startswith('Sessão futura')),'portalComprasPublicas':len(pcprows),'portalComprasPublicasForaDoPNCP':sum(1 for v in pcprows if str(v[-1]).startswith('Não')),'planilha':str(filepath)}
    (DATA/'estatisticas_finais.json').write_text(json.dumps(stats,ensure_ascii=False,indent=2))
    print(json.dumps(stats,ensure_ascii=False,indent=2))
    # Read back artifact structure and formulas.
    check=load_workbook(filepath,read_only=True,data_only=False)
    assert check['Itens para cotar'].max_row==len(itemrows)+1
    assert all(check[n].max_row==len(buckets[n])+1 for n in buckets)
    print('VALIDACAO OK',filepath)

if __name__=='__main__':main()
